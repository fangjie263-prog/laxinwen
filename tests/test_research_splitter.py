"""切割测试：HTML 有效 / PDF 有效 / DOCX 有效（对应需求二十六、二十七、三十七）。"""

from __future__ import annotations

import pytest

from news.research.sanitize import MAX_ARCHIVE_FILE_SIZE, split_extension
from news.research.splitter import (
    PYPDF_AVAILABLE,
    PYTHON_DOCX_AVAILABLE,
    part_stem,
    split_docx,
    split_file,
    split_html,
    verify_part,
)

from research_helpers import make_docx, make_html, make_pdf, write_bytes

requires_pypdf = pytest.mark.skipif(not PYPDF_AVAILABLE, reason="需要 pypdf")
requires_docx = pytest.mark.skipif(not PYTHON_DOCX_AVAILABLE, reason="需要 python-docx")


# ---------- HTML ----------

def test_split_html_produces_valid_parts_when_oversized(tmp_path):
    source = make_html(tmp_path / "big.html", sections=40, body="x" * 200_000, title="天齐锂业研究")
    assert source.stat().st_size > MAX_ARCHIVE_FILE_SIZE
    result = split_html(source, tmp_path / "out", stem="20260912A_Claude_09696.HK_天齐锂业_研究")
    assert result.ok, result.error
    assert len(result.parts) >= 2
    names = [path.name for path in result.parts]
    assert names[0].endswith("_Part01.html")
    assert names[1].endswith("_Part02.html")
    for path in result.parts:
        assert path.stat().st_size <= MAX_ARCHIVE_FILE_SIZE
        assert verify_part(path, ".html")
        text = path.read_text(encoding="utf-8")
        assert text.startswith("<!DOCTYPE html>")
        assert "</body>" in text


def test_split_html_small_file_not_split(tmp_path):
    source = make_html(tmp_path / "small.html", sections=2)
    result = split_html(source, tmp_path / "out", stem="研究")
    assert result.ok
    assert len(result.parts) == 1


def test_split_html_reports_oversized_single_block(tmp_path):
    source = make_html(tmp_path / "huge.html", sections=1, body="y" * (MAX_ARCHIVE_FILE_SIZE + 10))
    result = split_html(source, tmp_path / "out", stem="研究")
    assert not result.ok
    assert result.oversized
    assert "无法在大小限制内切割" in result.error


# ---------- PDF ----------

@requires_pypdf
def test_split_pdf_by_pages(tmp_path):
    # 每页写入大量元数据把文件撑大，迫使按页切割
    source = make_pdf(tmp_path / "big.pdf", pages=6, filler=40_000)
    max_bytes = max(1024, source.stat().st_size // 3)
    result = split_file(source, tmp_path / "out", stem="20260912A_Claude_09696.HK_天齐锂业_研究", max_bytes=max_bytes)
    assert result.ok, result.error
    assert len(result.parts) >= 2
    names = [path.name for path in result.parts]
    assert names[0].endswith("_Part01.pdf")
    for path in result.parts:
        assert path.stat().st_size <= max_bytes
        assert verify_part(path, ".pdf")
        assert path.read_bytes().startswith(b"%PDF")


@requires_pypdf
def test_split_pdf_single_page_over_limit_is_reported_not_corrupted(tmp_path):
    source = make_pdf(tmp_path / "one.pdf", pages=1, filler=40_000)
    limit = 5_000
    result = split_file(source, tmp_path / "out", stem="研究", max_bytes=limit)
    assert not result.ok
    assert result.oversized
    assert "无法安全切割" in result.error
    # 不允许留下任何超限的坏文件
    assert all(not path.exists() for path in (tmp_path / "out").glob("*") if path.suffix == ".pdf")
    assert source.is_file()  # 原文件仍在


@requires_pypdf
def test_split_pdf_keeps_all_pages_across_parts(tmp_path):
    from pypdf import PdfReader

    source = make_pdf(tmp_path / "multi.pdf", pages=8, filler=20_000)
    max_bytes = max(1024, source.stat().st_size // 4)
    result = split_file(source, tmp_path / "out", stem="研究", max_bytes=max_bytes)
    assert result.ok, result.error
    total_pages = sum(len(PdfReader(str(path)).pages) for path in result.parts)
    assert total_pages == 8


# ---------- DOCX ----------

@requires_docx
def test_split_docx_produces_valid_docx_parts(tmp_path):
    from docx import Document

    source = make_docx(tmp_path / "big.docx", chapters=120, paragraphs=60, text="研" * 400)
    # 分片上限必须大于 DOCX 容器开销（约 40KB），否则任何分片都无法合法生成
    max_bytes = max(50_000, source.stat().st_size // 3)
    result = split_docx(source, tmp_path / "out", stem="20260912A_Claude_09696.HK_天齐锂业_深度研究", max_bytes=max_bytes)
    assert result.ok, result.error
    assert len(result.parts) >= 2
    names = [path.name for path in result.parts]
    assert names[0].endswith("_Part01.docx")
    for path in result.parts:
        assert path.stat().st_size <= max_bytes
        assert verify_part(path, ".docx")
        # 每片都是合法 DOCX（可被 python-docx 打开且有内容）
        document = Document(str(path))
        assert len(document.paragraphs) > 0
        assert len(document.sections) >= 1


@requires_docx
def test_split_docx_does_not_duplicate_content(tmp_path):
    from docx import Document

    source = make_docx(tmp_path / "a.docx", chapters=200, paragraphs=30, text="内容" * 400)
    max_bytes = max(50_000, source.stat().st_size // 4)
    result = split_docx(source, tmp_path / "out", stem="研究", max_bytes=max_bytes)
    assert result.ok, result.error
    total_chars = sum(
        len(para.text) for path in result.parts for para in Document(str(path)).paragraphs
    )
    original_chars = sum(len(para.text) for para in Document(str(source)).paragraphs)
    # 完全一致：既不能丢内容，也不能重复内容（过去的实现会重复 sectPr/内容）
    assert total_chars == original_chars


@requires_docx
def test_split_docx_limit_below_container_reports_error(tmp_path):
    source = make_docx(tmp_path / "b.docx", chapters=5, paragraphs=5)
    result = split_docx(source, tmp_path / "out", stem="研究", max_bytes=1000)
    assert not result.ok
    assert "容器最小开销" in result.error


# ---------- 统一入口 ----------

def test_split_file_rejects_unknown_extension(tmp_path):
    source = write_bytes(tmp_path / "a.bin", 1024)
    result = split_file(source, tmp_path / "out", stem="研究")
    assert not result.ok
    assert "不支持切割的文件类型" in result.error


def test_part_stem_strips_existing_part_suffix():
    assert part_stem("20260912A_研究_Part01") == "20260912A_研究"
    assert part_stem("20260912A_研究_part12") == "20260912A_研究"
    assert part_stem("20260912A_研究") == "20260912A_研究"
    assert split_extension("a_Part01.pdf")[1] == ".pdf"
